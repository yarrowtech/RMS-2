"""Vendor ledger for the retailer Finance workspace.

One running account per vendor: every bill adds to what the retailer owes,
every payment / advance / debit note reduces it, with an opening balance, a
running balance after each line and a closing balance.

Additive by design: purchase invoices stay the source of truth for bills and
their payments, and cash payments are still recorded by the existing
purchase-invoice payment routine (so vendor notification, the finance-voucher
mirror, KYB checks and PO status updates behave exactly as before). The only
new stored data is the vendor *advance* (money paid on account, not yet set
against a bill) in its own collection.
"""
import re
from datetime import datetime
from html import escape
from typing import Any, List, Optional
from urllib.parse import unquote

from bson import ObjectId
from fastapi import APIRouter, Depends, HTTPException, status
from pydantic import BaseModel, Field

from ..db import finance_vouchers_collection, purchase_invoice_collection, vendor_advances_collection
from .finance_routes import aging_bucket, date_value, get_finance_context, number, record_scope
from .Purchaseinvoice_routes import PAYMENT_MODES, PaymentModel, _invoice_scope, record_payment

router = APIRouter(prefix="/api/finance/vendor-ledger", tags=["Retailer Finance - Vendor Ledger"])

# Bills the retailer actually owes: approved onwards. Drafts, held, rejected
# and cancelled invoices are not liabilities yet.
LIABLE_STATUSES = ["Approved", "PartiallyPaid", "Paid"]
PAYABLE_STATUSES = ["Approved", "PartiallyPaid"]


def _vendor_name(key: str) -> str:
    return unquote(key).strip()


def _name_query(name: str) -> dict:
    return {"$regex": f"^{re.escape(name)}$", "$options": "i"}


def _inv_vendor(inv: dict) -> str:
    return (inv.get("vendorName") or inv.get("vendor") or "Vendor").strip()


def _voucher_query(ctx: dict, name: str) -> dict:
    # Only debit notes raised against the vendor. Manual payment vouchers are
    # deliberately not counted: a payment is recorded against its bill (or as an
    # advance here), and counting a separate voucher too would count it twice.
    query = {
        **record_scope(ctx),
        "counterparty": _name_query(name),
        "voucher_type": "debit_note",
        "$or": [{"linked_invoice_id": ""}, {"linked_invoice_id": {"$exists": False}}, {"linked_invoice_id": None}],
    }
    return query


async def _load(ctx: dict, name: str) -> dict:
    """Everything that belongs to one vendor's account."""
    scope = record_scope(ctx)
    invoices = await purchase_invoice_collection.find(
        {**scope, "vendorName": _name_query(name), "status": {"$in": LIABLE_STATUSES}}
    ).to_list(length=5000)
    advances = await vendor_advances_collection.find({**scope, "vendor_name": _name_query(name)}).to_list(length=2000)
    vouchers = await finance_vouchers_collection.find(_voucher_query(ctx, name)).to_list(length=2000)
    return {"invoices": invoices, "advances": advances, "vouchers": vouchers}


def _lines(data: dict) -> list:
    """Every ledger event as {date, sort, type, ref, description, bill, credit, effect_note}.
    bill  = increases what is owed; credit = reduces it. Advance adjustments
    are shown but have no effect on the balance (the cash left when the
    advance was paid)."""
    rows = []
    for inv in data["invoices"]:
        ref = inv.get("invoiceNo") or "—"
        vinv = inv.get("vendorInvoiceNo") or ""
        rows.append({
            "date": date_value(inv.get("invoiceDate")), "sort": 0, "type": "bill", "ref": ref,
            "description": f"Bill {ref}" + (f" (vendor invoice {vinv})" if vinv else ""),
            "bill": number(inv.get("invoiceTotal")), "credit": 0.0, "invoice_id": str(inv["_id"]),
            "due_date": date_value(inv.get("dueDate")),
        })
        itemised = sum(number(p.get("amount")) for p in (inv.get("payments") or []))
        untracked = round(number(inv.get("paidAmount")) - itemised, 2)
        if untracked > 0.01:   # paid amount with no itemised payment (older / imported data)
            rows.append({
                "date": date_value(inv.get("invoiceDate")), "sort": 1, "type": "payment", "ref": ref,
                "description": f"Payment recorded on {ref} (not itemised)", "bill": 0.0, "credit": untracked, "invoice_id": str(inv["_id"]),
            })
        for pay in inv.get("payments", []) or []:
            adjusted = bool(pay.get("advanceAdjustment"))
            rows.append({
                "date": date_value(pay.get("paymentDate")), "sort": 1, "type": "advance_adjustment" if adjusted else "payment",
                "ref": pay.get("referenceNo") or ref,
                "description": (f"Advance set against {ref}" if adjusted else f"Payment against {ref} ({pay.get('paymentMode', '')})"),
                "bill": 0.0, "credit": 0.0 if adjusted else number(pay.get("amount")),
                "applied": number(pay.get("amount")) if adjusted else 0.0,
                "invoice_id": str(inv["_id"]),
            })
    for adv in data["advances"]:
        rows.append({
            "date": date_value(adv.get("date")), "sort": 1, "type": "advance", "ref": adv.get("reference_no") or "Advance",
            "description": f"Advance paid on account ({adv.get('payment_mode', '')})" + (f" - {adv.get('remarks')}" if adv.get("remarks") else ""),
            "bill": 0.0, "credit": number(adv.get("amount")),
        })
    for v in data["vouchers"]:
        note = v.get("voucher_type") == "debit_note"
        rows.append({
            "date": date_value(v.get("voucher_date")), "sort": 1, "type": "debit_note" if note else "voucher_payment",
            "ref": v.get("voucher_no") or v.get("id") or "",
            "description": (v.get("remarks") or ("Debit note" if note else "Payment voucher")),
            "bill": 0.0, "credit": number(v.get("amount")),
        })
    rows.sort(key=lambda r: (r["date"], r["sort"]))
    return rows


def _summary(data: dict) -> dict:
    today = datetime.utcnow()
    today_s = today.strftime("%Y-%m-%d")
    billed = paid_on_bills = gross_due = overdue = 0.0
    advance_applied = 0.0
    open_bills = []
    for inv in data["invoices"]:
        billed += number(inv.get("invoiceTotal"))
        paid_on_bills += number(inv.get("paidAmount"))
        for pay in inv.get("payments", []) or []:
            if pay.get("advanceAdjustment"):
                advance_applied += number(pay.get("amount"))
        balance = number(inv.get("balanceDue"))
        if balance > 0 and inv.get("status") in PAYABLE_STATUSES:
            gross_due += balance
            due = date_value(inv.get("dueDate"))
            late = 0
            if due and due < today_s:
                overdue += balance
                late = (today.date() - datetime.strptime(due, "%Y-%m-%d").date()).days
            open_bills.append({
                "id": str(inv["_id"]), "invoice_no": inv.get("invoiceNo") or "—", "invoice_date": date_value(inv.get("invoiceDate")),
                "due_date": due, "invoice_total": number(inv.get("invoiceTotal")), "paid_amount": number(inv.get("paidAmount")),
                "balance_due": balance, "days_overdue": late, "aging_bucket": aging_bucket(due, today),
            })
    open_bills.sort(key=lambda b: (b["due_date"] or "9999-12-31", b["invoice_date"]))
    advances_paid = sum(number(a.get("amount")) for a in data["advances"])
    advance_available = round(max(0.0, advances_paid - advance_applied), 2)
    credits = sum(number(v.get("amount")) for v in data["vouchers"])
    cash_on_bills = round(paid_on_bills - advance_applied, 2)
    outstanding = round(billed - cash_on_bills - advances_paid - credits, 2)
    return {
        "total_billed": round(billed, 2), "paid_on_bills": round(paid_on_bills, 2), "advance_paid": round(advances_paid, 2),
        "advance_available": advance_available, "credits_and_notes": round(credits, 2), "gross_bills_due": round(gross_due, 2),
        "outstanding": outstanding, "overdue": round(overdue, 2), "open_bills": open_bills, "invoice_count": len(data["invoices"]),
    }


@router.get("/vendors")
async def vendor_summary(ctx: dict = Depends(get_finance_context)):
    scope = record_scope(ctx)
    names = set()
    async for inv in purchase_invoice_collection.find({**scope, "status": {"$in": LIABLE_STATUSES}}, {"vendorName": 1, "vendor": 1}):
        names.add(_inv_vendor(inv))
    async for adv in vendor_advances_collection.find(scope, {"vendor_name": 1}):
        if adv.get("vendor_name"):
            names.add(adv["vendor_name"].strip())
    rows = []
    for name in sorted(names, key=str.lower):
        data = await _load(ctx, name)
        s = _summary(data)
        last = max((l["date"] for l in _lines(data) if l["type"] in ("payment", "advance")), default="")
        rows.append({
            "vendor": name, "vendor_key": name, "invoice_count": s["invoice_count"], "total_billed": s["total_billed"],
            "paid": round(s["paid_on_bills"] + s["credits_and_notes"], 2), "advance_available": s["advance_available"],
            "outstanding": s["outstanding"], "overdue": s["overdue"], "last_payment": last,
        })
    rows.sort(key=lambda r: -r["outstanding"])
    totals = {
        "vendors": len(rows), "total_billed": round(sum(r["total_billed"] for r in rows), 2),
        "outstanding": round(sum(r["outstanding"] for r in rows), 2), "overdue": round(sum(r["overdue"] for r in rows), 2),
        "advance_available": round(sum(r["advance_available"] for r in rows), 2),
    }
    return {"rows": rows, "totals": totals}


async def _compute_statement(ctx: dict, vendor_key: str, from_date: Optional[str], to_date: Optional[str]) -> dict:
    """Shared by the JSON statement route and the PDF/Excel download route,
    so the downloaded file can never drift out of sync with what the screen
    shows — both read off this exact same computation."""
    name = _vendor_name(vendor_key)
    if not name:
        raise HTTPException(status_code=400, detail="Vendor is required.")
    data = await _load(ctx, name)
    if not data["invoices"] and not data["advances"] and not data["vouchers"]:
        raise HTTPException(status_code=404, detail="No ledger entries found for this vendor.")
    lines = _lines(data)
    opening = 0.0
    shown = []
    running = 0.0
    for line in lines:
        running = round(running + line["bill"] - line["credit"], 2)
        line["balance"] = running
        if from_date and line["date"] < from_date:
            opening = running
            continue
        if to_date and line["date"] > to_date:
            continue
        shown.append(line)
    closing = shown[-1]["balance"] if shown else opening
    summary = _summary(data)
    return {
        "vendor": name, "from_date": from_date or "", "to_date": to_date or "", "opening_balance": round(opening, 2),
        "closing_balance": round(closing, 2), "lines": shown, "summary": {k: v for k, v in summary.items() if k != "open_bills"},
        "open_bills": summary["open_bills"],
    }


@router.get("/vendors/{vendor_key}/statement")
async def vendor_statement(vendor_key: str, from_date: Optional[str] = None, to_date: Optional[str] = None,
                           ctx: dict = Depends(get_finance_context)):
    return await _compute_statement(ctx, vendor_key, from_date, to_date)


class Allocation(BaseModel):
    invoice_id: str
    amount: float = Field(..., gt=0)


class VendorPayment(BaseModel):
    amount: float = Field(..., gt=0)
    paymentDate: str = Field(default_factory=lambda: datetime.utcnow().strftime("%Y-%m-%d"))
    paymentMode: str = "NEFT"
    referenceNo: str = ""
    remarks: str = ""
    allocations: Optional[List[Allocation]] = None   # None = oldest due first
    keep_excess_as_advance: bool = True


class AdvanceCreate(BaseModel):
    amount: float = Field(..., gt=0)
    date: str = Field(default_factory=lambda: datetime.utcnow().strftime("%Y-%m-%d"))
    payment_mode: str = "NEFT"
    reference_no: str = ""
    remarks: str = ""


async def _new_advance(ctx: dict, name: str, amount: float, date: str, mode: str, ref: str, remarks: str) -> dict:
    if mode not in PAYMENT_MODES:
        raise HTTPException(status_code=400, detail=f"Invalid payment mode '{mode}'.")
    try:
        datetime.strptime(date, "%Y-%m-%d")
    except ValueError:
        raise HTTPException(status_code=400, detail="Date must use YYYY-MM-DD.")
    now = datetime.utcnow()
    scope = record_scope(ctx)
    doc = {
        "id": str(ObjectId()), **scope, "vendor_name": name, "amount": round(amount, 2), "date": date, "payment_mode": mode,
        "reference_no": ref.strip(), "remarks": remarks.strip(), "created_by": ctx.get("admin_id", ""), "created_at": now,
    }
    await vendor_advances_collection.insert_one(doc)
    # Cash left the business now, so it is mirrored into the finance vouchers like any other payment.
    await finance_vouchers_collection.insert_one({
        "id": doc["id"], "voucher_no": f"AP-ADV-{doc['id'][-8:].upper()}", **scope, "voucher_type": "payment", "voucher_date": date,
        "amount": doc["amount"], "account": "Vendor advance", "counterparty": name, "payment_mode": mode, "reference_no": ref.strip(),
        "remarks": remarks.strip() or "Advance paid on account", "category": "vendor_advance", "linked_invoice_id": "advance:" + doc["id"],
        "linked_sale_id": "", "source": "vendor_ledger", "status": "Posted", "created_by": ctx.get("admin_id", ""), "created_at": now,
    })
    doc.pop("_id", None)
    return doc


@router.post("/vendors/{vendor_key}/advance", status_code=status.HTTP_201_CREATED)
async def record_advance(vendor_key: str, payload: AdvanceCreate, ctx: dict = Depends(get_finance_context)):
    name = _vendor_name(vendor_key)
    if not name:
        raise HTTPException(status_code=400, detail="Vendor is required.")
    doc = await _new_advance(ctx, name, payload.amount, payload.date, payload.payment_mode, payload.reference_no, payload.remarks)
    return {"message": f"Advance of ₹{doc['amount']:,.2f} recorded for {name}.", "advance": doc}


@router.post("/vendors/{vendor_key}/pay")
async def pay_vendor(vendor_key: str, payload: VendorPayment, ctx: dict = Depends(get_finance_context)):
    """One payment to a vendor, spread over their open bills (oldest due first,
    or the split you choose). Anything left over becomes an advance."""
    name = _vendor_name(vendor_key)
    if payload.paymentMode not in PAYMENT_MODES:
        raise HTTPException(status_code=400, detail=f"Invalid payment mode '{payload.paymentMode}'.")
    data = await _load(ctx, name)
    open_bills = {b["id"]: b for b in _summary(data)["open_bills"]}

    if payload.allocations is None:
        plan, left = [], round(payload.amount, 2)
        for bill in open_bills.values():
            if left <= 0.004:
                break
            take = round(min(left, bill["balance_due"]), 2)
            plan.append((bill["id"], take))
            left = round(left - take, 2)
    else:
        merged: dict = {}
        for a in payload.allocations:
            if a.invoice_id not in open_bills:
                raise HTTPException(status_code=400, detail="One of the selected bills is not an open bill of this vendor.")
            merged[a.invoice_id] = round(merged.get(a.invoice_id, 0.0) + a.amount, 2)
        plan = []
        for invoice_id, amount in merged.items():
            bill = open_bills[invoice_id]
            if amount > bill["balance_due"] + 0.01:
                raise HTTPException(status_code=400, detail=f"₹{amount:,.2f} is more than the ₹{bill['balance_due']:,.2f} due on {bill['invoice_no']}.")
            plan.append((invoice_id, amount))
    allocated = round(sum(amount for _, amount in plan), 2)
    if allocated > payload.amount + 0.01:
        raise HTTPException(status_code=400, detail="The split is more than the amount being paid.")
    excess = round(payload.amount - allocated, 2)
    if excess > 0.01 and not payload.keep_excess_as_advance:
        raise HTTPException(status_code=400, detail=f"₹{excess:,.2f} is left over. Allocate it to a bill or keep it as an advance.")

    # The purchase-invoice payment routine is tenant-scoped through a request-scoped
    # context value that its own router normally sets - set it the same way here.
    _invoice_scope.set({"kind": "admin", "tenant_id": ctx["tenant_id"], "store_id": ctx.get("store_id")})
    applied = []
    for invoice_id, amount in plan:
        if amount <= 0:
            continue
        result = await record_payment(invoice_id, PaymentModel(
            amount=amount, paymentDate=payload.paymentDate, paymentMode=payload.paymentMode, referenceNo=payload.referenceNo,
            remarks=payload.remarks or f"Part of one payment to {name}",
        ))
        applied.append({"invoice_id": invoice_id, "invoice_no": open_bills[invoice_id]["invoice_no"], "amount": amount, "balance_due": result.get("balanceDue")})

    advance = None
    if excess > 0.01:
        advance = await _new_advance(ctx, name, excess, payload.paymentDate, payload.paymentMode, payload.referenceNo, payload.remarks or "Excess of payment kept as advance")
    return {"message": f"Paid ₹{payload.amount:,.2f} to {name}.", "applied": applied, "advance_created": advance}


@router.post("/vendors/{vendor_key}/apply-advance")
async def apply_advance(vendor_key: str, ctx: dict = Depends(get_finance_context)):
    """Sets the vendor's unused advance against their open bills, oldest due first.
    No new cash moves - the advance was already paid - so no voucher is created."""
    name = _vendor_name(vendor_key)
    data = await _load(ctx, name)
    summary = _summary(data)
    left = summary["advance_available"]
    if left <= 0.004:
        raise HTTPException(status_code=400, detail="This vendor has no unused advance.")
    if not summary["open_bills"]:
        raise HTTPException(status_code=400, detail="This vendor has no open bills to set the advance against.")
    applied = []
    today = datetime.utcnow().strftime("%Y-%m-%d")
    for bill in summary["open_bills"]:
        if left <= 0.004:
            break
        take = round(min(left, bill["balance_due"]), 2)
        doc = await purchase_invoice_collection.find_one({"_id": ObjectId(bill["id"])})
        new_paid = round(number(doc.get("paidAmount")) + take, 2)
        new_balance = round(max(0.0, number(doc.get("invoiceTotal")) - new_paid), 2)
        await purchase_invoice_collection.update_one(
            {"_id": doc["_id"]},
            {"$set": {"paidAmount": new_paid, "balanceDue": new_balance, "advanceAdjusted": round(number(doc.get("advanceAdjusted")) + take, 2),
                      "paymentStatus": "Paid" if new_balance < 0.01 else "Partial", "status": "Paid" if new_balance < 0.01 else "PartiallyPaid",
                      "updatedAt": datetime.utcnow()},
             "$push": {"payments": {"paymentId": str(ObjectId()), "amount": take, "paymentDate": today, "paymentMode": "Advance",
                                    "referenceNo": "ADVANCE-ADJ", "remarks": "Advance set against this bill", "isAdvance": True,
                                    "advanceAdjustment": True, "scheduleId": "", "pdcPending": False, "recordedAt": datetime.utcnow().isoformat()}}},
        )
        applied.append({"invoice_no": bill["invoice_no"], "amount": take, "balance_due": new_balance})
        left = round(left - take, 2)
    return {"message": f"Set ₹{summary['advance_available'] - left:,.2f} of advance against {len(applied)} bill(s).", "applied": applied}


# ─────────────────────────────────────────────
# DOWNLOAD STATEMENT (PDF / Excel) — same two-format pattern as the vendor
# PO export (vendor_routes.py) and the GRC export (grc_routes.py). Built off
# _compute_statement so the file always matches the on-screen statement.
# ─────────────────────────────────────────────

_LINE_TYPE_LABEL = {
    "bill": "Bill", "payment": "Payment", "advance": "Advance paid", "advance_adjustment": "Advance set against bill",
    "debit_note": "Debit note", "voucher_payment": "Payment voucher",
}


def _statement_export_response(statement: dict, fmt: str):
    from fastapi.responses import StreamingResponse
    from io import BytesIO

    vendor = statement["vendor"]
    safe_name = re.sub(r"[^A-Za-z0-9]+", "-", vendor).strip("-") or "vendor"
    period = f"{statement['from_date'] or 'inception'} to {statement['to_date'] or 'date'}"
    s = statement["summary"]
    open_bills = statement["open_bills"]
    lines = statement["lines"]

    if fmt == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        wb = Workbook(); ws = wb.active; ws.title = "Statement"
        ws.append(["VENDOR LEDGER STATEMENT", vendor])
        ws.append(["Period", period])
        ws.append(["Total billed", s["total_billed"]]); ws.append(["Paid on bills", s["paid_on_bills"]])
        ws.append(["Advance available", s["advance_available"]]); ws.append(["Outstanding", s["outstanding"]]); ws.append(["Overdue", s["overdue"]])
        ws.append([])

        if open_bills:
            ws.append(["OPEN BILLS"])
            ws.append(["Bill", "Bill date", "Due date", "Total", "Paid", "Balance due", "Days overdue"])
            header_row = ws.max_row
            for cell in ws[header_row]:
                cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="0EA5E9")
            for b in open_bills:
                ws.append([b["invoice_no"], b["invoice_date"], b["due_date"] or "", b["invoice_total"], b["paid_amount"], b["balance_due"], b["days_overdue"]])
            ws.append([])

        ws.append(["STATEMENT"])
        ws.append(["Date", "Entry", "Reference", "Description", "Bill (+)", "Paid / credit (-)", "Balance"])
        header_row = ws.max_row
        for cell in ws[header_row]:
            cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="4F46E5"); cell.alignment = Alignment(horizontal="center")
        ws.append(["", "Opening balance", "", "", "", "", statement["opening_balance"]])
        ws.cell(ws.max_row, 2).font = Font(bold=True)
        for l in lines:
            ws.append([
                l["date"], _LINE_TYPE_LABEL.get(l["type"], l["type"]), l.get("ref") or "", l.get("description") or "",
                l["bill"] or "", (l.get("applied") if l["type"] == "advance_adjustment" else l["credit"]) or "", l["balance"],
            ])
        ws.append(["", "Closing balance", "", "", "", "", statement["closing_balance"]])
        ws.cell(ws.max_row, 2).font = Font(bold=True); ws.cell(ws.max_row, 7).font = Font(bold=True)
        for column in ws.columns:
            ws.column_dimensions[column[0].column_letter].width = min(max(max(len(str(cell.value or "")) for cell in column) + 3, 12), 48)
        buf = BytesIO(); wb.save(buf); buf.seek(0)
        return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="{safe_name}-ledger.xlsx"'})

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=15 * mm, rightMargin=15 * mm, topMargin=16 * mm, bottomMargin=16 * mm)
    styles = getSampleStyleSheet()
    title = ParagraphStyle("VendorLedgerTitle", parent=styles["Heading1"], textColor=colors.HexColor("#4F46E5"))
    elements = [
        Paragraph("Vendor Ledger Statement", title),
        Paragraph(
            f"Vendor: {vendor}<br/>Period: {period}<br/>"
            f"Total billed: {s['total_billed']:,.2f} &nbsp;|&nbsp; Paid: {s['paid_on_bills']:,.2f} &nbsp;|&nbsp; "
            f"Advance available: {s['advance_available']:,.2f} &nbsp;|&nbsp; Outstanding: {s['outstanding']:,.2f} &nbsp;|&nbsp; Overdue: {s['overdue']:,.2f}",
            styles["Normal"],
        ),
        Spacer(1, 6 * mm),
    ]

    if open_bills:
        elements.append(Paragraph("Open bills", styles["Heading3"]))
        ob_data = [["Bill", "Bill date", "Due date", "Total", "Paid", "Balance due", "Overdue"]]
        for b in open_bills:
            ob_data.append([
                b["invoice_no"], b["invoice_date"], b["due_date"] or "-", f"{b['invoice_total']:,.2f}",
                f"{b['paid_amount']:,.2f}", f"{b['balance_due']:,.2f}", f"{b['days_overdue']}d" if b["days_overdue"] else "-",
            ])
        ob_table = Table(ob_data, repeatRows=1, colWidths=[35 * mm, 28 * mm, 28 * mm, 32 * mm, 32 * mm, 32 * mm, 25 * mm])
        ob_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0EA5E9")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("GRID", (0, 0), (-1, -1), .35, colors.HexColor("#CBD5E1")),
            ("ALIGN", (3, 1), (-1, -1), "RIGHT"), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]))
        elements.extend([ob_table, Spacer(1, 8 * mm)])

    elements.append(Paragraph("Statement", styles["Heading3"]))
    st_data = [["Date", "Entry", "Reference", "Description", "Bill (+)", "Paid / credit (-)", "Balance"]]
    st_data.append(["", "Opening balance", "", "", "", "", f"{statement['opening_balance']:,.2f}"])
    for l in lines:
        credit_value = l.get("applied") if l["type"] == "advance_adjustment" else l["credit"]
        st_data.append([
            l["date"], _LINE_TYPE_LABEL.get(l["type"], l["type"]), Paragraph(escape(str(l.get("ref") or "")), styles["BodyText"]),
            Paragraph(escape(str(l.get("description") or "")), styles["BodyText"]),
            f"{l['bill']:,.2f}" if l["bill"] else "", f"{credit_value:,.2f}" if credit_value else "", f"{l['balance']:,.2f}",
        ])
    st_data.append(["", "Closing balance", "", "", "", "", f"{statement['closing_balance']:,.2f}"])
    st_table = Table(st_data, repeatRows=1, colWidths=[22 * mm, 30 * mm, 28 * mm, 90 * mm, 24 * mm, 28 * mm, 28 * mm])
    st_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4F46E5")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTNAME", (0, 1), (-1, 1), "Helvetica-Bold"), ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), .35, colors.HexColor("#CBD5E1")), ("ALIGN", (4, 1), (-1, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    elements.extend([st_table, Spacer(1, 6 * mm), Paragraph(f"Generated from RMS on {datetime.utcnow().strftime('%d %b %Y, %H:%M UTC')}.", styles["Normal"])])
    doc.build(elements); buf.seek(0)
    return StreamingResponse(buf, media_type="application/pdf", headers={"Content-Disposition": f'attachment; filename="{safe_name}-ledger.pdf"'})


@router.get("/vendors/{vendor_key}/statement/download")
async def download_vendor_statement(vendor_key: str, format: str = "pdf", from_date: Optional[str] = None,
                                    to_date: Optional[str] = None, ctx: dict = Depends(get_finance_context)):
    """Download a vendor's ledger statement as a printable PDF or an editable Excel sheet."""
    fmt = format.lower()
    if fmt not in {"pdf", "xlsx"}:
        raise HTTPException(status_code=400, detail="format must be pdf or xlsx")
    statement = await _compute_statement(ctx, vendor_key, from_date, to_date)
    return _statement_export_response(statement, fmt)
